"""End-to-end smoke tests on a synthetic input directory (no external data)."""

import os
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from src.cli import main
from src.pipeline.export import FINAL_COLUMNS
from src.pipeline.inputs import InputError, discover_inputs
from src.pipeline.report import RunReport
from src.pipeline.run import run_pipeline, write_outputs
from src.pipeline.settings import PipelineConfig, load_config
from tests.synthetic import make_synthetic_dirs, write_hub_names


@pytest.fixture
def dirs(tmp_path):
    return make_synthetic_dirs(tmp_path)


def test_run_pipeline_end_to_end(dirs, tmp_path):
    input_dir, ref_dir = dirs
    cfg = load_config(None, {"mc_iterations": 300, "apply_eligibility_filter": False, "extra_columns": "hub_id,eligible", "keep_intermediates": True})
    inputs = discover_inputs(input_dir, ref_dir)
    report = RunReport()
    result = run_pipeline(cfg, inputs, report)

    res = result.results
    assert len(res) == 3  # hub A (nodes 1,2,3 merged: 1+2 by distance, 3 by is_same_group), hub B, lone rail node
    a = res[res["node"].map(lambda ns: 1 in ns)].iloc[0]
    b = res[res["node"].map(lambda ns: 4 in ns)].iloc[0]
    assert set(a["node"]) == {1, 2, 3} and a["TotalDemand"] == 120000 and a["TotalTransfers"] == 10000
    assert a["area"] == "תל אביב" and a["location"] == ["גלעין"] and a["Metro"] == "תל אביב"
    assert a["HubType"] == "מטרופוליני" and a["bus_terminal"] == 3  # suburban rail + metro + LRT, >= 5000
    assert a["pop_0_500"] > 0 and a["TotalEmp_2050"] > 0
    assert b["TotalDemand"] == 2500  # manual override (transfers kept)
    assert b["NumLinesStatus_6"] == 1 and b["NumLinesStatus_0"] == 0
    assert 'אדום צפון (רק"ל)' in a["Line_Names_forPlot"] and "מטרו צפון (מטרו)" in a["Line_Names_forPlot"]
    assert set(res["RankByHubTypeMetro"].dropna()) <= {1, 2}
    assert res["eligible"].sum() == 2  # the rail-only lone node is not eligible

    warnings = " | ".join(w.message for w in report.warnings)
    assert "no row in the lines/mode file" in warnings and "duplicated Line_ModelName" in warnings

    paths = write_outputs(result, tmp_path / "out")
    wb = load_workbook(paths["xlsx"])
    ws = wb[cfg.output_sheet_name]
    assert list(ws.tables) == ["טבלה1"]
    header = [c.value for c in ws[1]]
    assert header == list(FINAL_COLUMNS) + ["hub_id", "eligible"]
    assert (tmp_path / "out" / "intermediate" / "groups.geojson").exists()
    assert (tmp_path / "out" / "hub_identity.csv").exists()
    assert "Metrics" in paths["report_md"].read_text(encoding="utf-8")
    # the shareable H3 layer: hub cells, their catchment, base attributes everywhere
    import geopandas as gpd

    layer = gpd.read_file(paths["h3_layer"], layer="h3_cells")
    assert paths["h3_layer"].name == "h3_layer.gpkg" and layer.crs.to_epsg() == 2039
    hub_rows = layer[layer["role"] == "hub"]
    assert len(hub_rows) == len(result.hexes) and set(hub_rows["group"]) == set(res["group"])
    assert (layer["role"] == "influence").sum() > 100 and (layer["pop_2050"] > 0).any()
    hub_a_cells = hub_rows[hub_rows["group"] == int(a["group"])]
    assert set(";".join(hub_a_cells["nodes"]).split(";")) == {"1", "2", "3"}
    assert (hub_a_cells["HubType"] == "מטרופוליני").all() and hub_a_cells["scored"].all()
    assert hub_a_cells["TotalScore_MC"].notna().all()
    infl = layer[layer["role"] == "influence"]
    assert (infl["dist_nearest_hub_m"] <= 1500).all() and infl["hubs_within"].notna().all()


def test_rerun_is_byte_identical(dirs, tmp_path):
    input_dir, ref_dir = dirs
    cfg = load_config(None, {"mc_iterations": 200})
    inputs = discover_inputs(input_dir, ref_dir)
    first = write_outputs(run_pipeline(cfg, inputs), tmp_path / "a")["csv"].read_bytes()
    second = write_outputs(run_pipeline(cfg, inputs), tmp_path / "b")["csv"].read_bytes()
    assert first == second


def test_missing_required_layer_fails_loudly(dirs):
    input_dir, ref_dir = dirs
    for ext in (".shp", ".shx", ".dbf", ".prj", ".cpg"):
        p = ref_dir / f"TAZ_1270{ext}"
        if p.exists():
            p.unlink()
    inputs = discover_inputs(input_dir, ref_dir)
    # the default path reads the base layer, so a missing TAZ shapefile does not matter ...
    assert run_pipeline(load_config(None, {"mc_iterations": 50}), inputs).results["pop_0_500"].gt(0).any()
    # ... but the shapefile path needs it
    with pytest.raises(InputError) as exc:
        run_pipeline(load_config(None, {"spatial_source": "shapefiles"}), inputs)
    assert any("'taz'" in p for p in exc.value.problems)
    # tolerated only when explicitly allowed (scores of 0)
    result = run_pipeline(load_config(None, {"mc_iterations": 50, "spatial_source": "shapefiles"}), inputs, allow_missing_layers=True)
    assert (result.results["pop_0_500"] == 0).all()


def test_missing_base_layer_fails_loudly(tmp_path):
    from tests.synthetic import make_synthetic_dirs as make

    input_dir, ref_dir = make(tmp_path, with_base_layer=False)
    inputs = discover_inputs(input_dir, ref_dir)
    with pytest.raises(InputError) as exc:
        run_pipeline(PipelineConfig(), inputs)
    assert any("'h3_base'" in p and "prepare-base" in p for p in exc.value.problems)
    # validation without a configured source: the shapefiles stand in for the base layer
    assert inputs.missing_required() == []
    assert [s.key for s in inputs.missing_required("h3_base")] == ["h3_base"]
    assert run_pipeline(load_config(None, {"mc_iterations": 50, "spatial_source": "shapefiles"}), inputs).results["pop_0_500"].gt(0).any()


def test_hub_names_by_h3(dirs, tmp_path):
    input_dir, ref_dir = dirs
    inputs = discover_inputs(input_dir, ref_dir)
    result = run_pipeline(load_config(None, {"mc_iterations": 50, "apply_eligibility_filter": False}), inputs)
    hexes = {int(r["group"]): list(r["h3_index"]) for _, r in result.results.iterrows()}
    a_group = int(result.results[result.results["node"].map(lambda ns: 1 in ns)]["group"].iloc[0])
    write_hub_names(ref_dir, hexes, {a_group: "מרכז סבידור"})
    again = run_pipeline(load_config(None, {"mc_iterations": 50, "apply_eligibility_filter": False}), discover_inputs(input_dir, ref_dir))
    assert again.results.set_index("group").loc[a_group, "HubNameHE"] == "מרכז סבידור"


def test_cli_run(dirs, tmp_path, capsys, monkeypatch):
    input_dir, ref_dir = dirs
    out = tmp_path / "cli_out"
    code = main(["run", "--input-dir", str(input_dir), "--reference-dir", str(ref_dir), "--output-dir", str(out), "--set", "mc_iterations=100"])
    assert code == 0
    assert (out / "hub_prioritization_results.xlsx").exists() and (out / "run.log").exists()
    assert "hubs written" in capsys.readouterr().out

    assert (out / "h3_layer.gpkg").exists()

    # the base layer can be exported on its own
    assert main(["export-h3", "--reference-dir", str(ref_dir), "--out", str(tmp_path / "cells"), "--format", "csv"]) == 0
    cells = pd.read_csv(tmp_path / "cells.csv")
    assert {"h3_index", "area", "pop_2050", "geometry_wkt"} <= set(cells.columns) and (cells["role"] == "base").all()

    # without the base layer the default run stops with a hint ...
    (ref_dir / "h3_base.parquet").unlink()
    assert main(["run", "--input-dir", str(input_dir), "--reference-dir", str(ref_dir), "--output-dir", str(out)]) == 1
    assert "prepare-base" in capsys.readouterr().out
    # ... the shapefile path still works, and a missing layer is a hard error for it ...
    shp = ["--set", "spatial_source=shapefiles", "--set", "mc_iterations=50"]
    assert main(["run", "--input-dir", str(input_dir), "--reference-dir", str(ref_dir), "--output-dir", str(out), *shp]) == 0
    for ext in (".shp", ".shx", ".dbf", ".prj", ".cpg"):
        (ref_dir / f"BUS_TERMINAL_STRAT{ext}").unlink(missing_ok=True)
    assert main(["run", "--input-dir", str(input_dir), "--reference-dir", str(ref_dir), "--output-dir", str(out), *shp]) == 1
    # ... unless the escape hatch is set
    monkeypatch.setenv("HUBS_ALLOW_MISSING_LAYERS", "1")
    assert main(["run", "--input-dir", str(input_dir), "--reference-dir", str(ref_dir), "--output-dir", str(out), *shp]) == 0

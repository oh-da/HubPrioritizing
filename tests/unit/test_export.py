import pandas as pd
import pytest
from openpyxl import load_workbook

from src.pipeline.export import FINAL_COLUMNS, read_results_xlsx, select_final_columns, write_results_csv, write_results_xlsx


def _frame(n=3):
    df = pd.DataFrame({c: [0] * n for c in FINAL_COLUMNS})
    df["group"] = list(range(n))
    df["HubNameHE"] = ["תחנת רכבת סבידור מרכז", "מרכזית המפרץ", "x"]
    df["node"] = [[1, 2], [3], [4]]
    df["Mode_Planned"] = [["Metro", "LRT"], ["BRT"], []]
    df["location"] = [["גלעין"], ["טבעת"], ["חוץ"]]
    df["h3_index"] = [["a"], ["b"], ["c"]]
    df["Line_Unique"] = [["L1"], ["L2"], []]
    df["hub_id"] = ["Ha", "Hb", "Hc"]
    return df


def test_select_final_columns_order_and_lists():
    out = select_final_columns(_frame(), extra_columns=["hub_id", "missing"])
    assert list(out.columns) == list(FINAL_COLUMNS) + ["hub_id"]
    assert out.loc[0, "node"] == "[1, 2]" and out.loc[0, "Mode_Planned"] == "['Metro', 'LRT']"
    with pytest.raises(ValueError):
        select_final_columns(_frame().drop(columns=["x"]))


def test_write_results_xlsx_table_and_hebrew(tmp_path):
    path = write_results_xlsx(_frame(), tmp_path / "out" / "hub_prioritization_results.xlsx", sheet_name="hubs_final_results", table_name="טבלה1")
    wb = load_workbook(path)
    ws = wb["hubs_final_results"]
    assert len(wb.sheetnames) == 1
    assert list(ws.tables.keys()) == ["טבלה1"]
    assert ws.tables["טבלה1"].ref == "A1:BR4"
    back = read_results_xlsx(path)
    assert list(back.columns) == list(FINAL_COLUMNS)
    assert back.loc[0, "HubNameHE"] == "תחנת רכבת סבידור מרכז"
    assert not any(str(v).startswith("=") for v in back.iloc[0].tolist())


def test_write_results_csv_round_trip(tmp_path):
    path = write_results_csv(_frame(), tmp_path / "r.csv")
    back = pd.read_csv(path, encoding="utf-8-sig")
    assert list(back.columns) == list(FINAL_COLUMNS) and len(back) == 3
